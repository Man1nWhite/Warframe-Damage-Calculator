"""SQLite connection context managers do not themselves close the connection."""
from __future__ import annotations

import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

import openpyxl

from modules.database import ModRepository, WeaponRepository
from modules.importer import ImportPayload, build_database, rows_from_spreadsheet

ROOT = Path(__file__).resolve().parents[1]
DATABASE = ROOT / "data" / "weapons.db"


class TrackingConnection(sqlite3.Connection):
    closed = False

    def close(self):
        self.closed = True
        super().close()


class ConnectionLifecycleTests(unittest.TestCase):
    def test_repository_queries_close_connections_on_success_and_error(self):
        connections = []
        def connect():
            connection = sqlite3.connect(DATABASE, factory=TrackingConnection)
            connection.row_factory = sqlite3.Row
            connections.append(connection)
            return connection

        weapons = WeaponRepository(DATABASE)
        mods = ModRepository(DATABASE)
        with patch.object(weapons, "connect", side_effect=connect), \
             patch.object(mods, "connect", side_effect=connect):
            self.assertTrue(weapons.search_weapons("Hek"))
            self.assertIsNotNone(weapons.get_weapon("Hek"))
            self.assertTrue(mods.search_mods("Serration", "PRIMARY"))
            available = mods.list_mods_for_slot("PRIMARY")
            self.assertTrue(available)
            self.assertIsNotNone(mods.get_mod(available[0].unique_name))
            mods.fetch_ranked_effects(available[0].unique_name, 0)
        self.assertEqual(len(connections), 6)
        self.assertTrue(all(connection.closed for connection in connections))

        def empty_connect():
            connection = sqlite3.connect(":memory:", factory=TrackingConnection)
            connections.append(connection)
            return connection
        with patch.object(weapons, "connect", side_effect=empty_connect):
            self.assertEqual(weapons.search_weapons("Hek"), [])
        self.assertTrue(connections[-1].closed)

    def test_uninitialised_weapon_catalog_is_empty_but_partial_schema_is_an_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "weapons.db"
            repository = WeaponRepository(path)
            self.assertEqual(repository.search_weapons(""), [])
            with closing(sqlite3.connect(path)) as connection:
                connection.execute("CREATE TABLE unrelated (value TEXT)")
                connection.commit()
            with self.assertRaisesRegex(sqlite3.OperationalError, "no such table: weapons"):
                repository.search_weapons("")

    def test_ranked_effect_batch_matches_individual_queries_and_uses_one_connection(self):
        repository = ModRepository(DATABASE)
        mods = repository.list_mods_for_slot("PRIMARY")
        requests = [(mod.unique_name, rank) for mod in mods[:8] for rank in (0, mod.fusion_limit)]
        requests.append(requests[0])  # duplicate slots share the same lookup
        expected = {key: repository.fetch_ranked_effects(*key) for key in requests}
        with patch.object(repository, "connect", wraps=repository.connect) as connect:
            actual = repository.fetch_ranked_effects_many(requests)
        self.assertEqual(actual, expected)
        self.assertEqual(connect.call_count, 1)
        with patch.object(repository, "connect", wraps=repository.connect) as connect:
            self.assertEqual(repository.fetch_ranked_effects_many([]), {})
        connect.assert_not_called()

    def test_ranked_effect_batch_chunks_large_iterables_and_reports_missing_tables(self):
        repository = ModRepository(DATABASE)
        mod = repository.list_mods_for_slot("PRIMARY")[0]
        real_key = (mod.unique_name, mod.fusion_limit)
        requests = [(f"/missing/{i}", 0) for i in range(801)]
        requests.insert(400, real_key)
        expected_effects = repository.fetch_ranked_effects(*real_key)
        with patch.object(repository, "connect", wraps=repository.connect) as connect:
            actual = repository.fetch_ranked_effects_many(iter(requests))
        self.assertEqual(connect.call_count, 1)
        self.assertEqual(len(actual), len(requests))
        self.assertEqual(actual[real_key], expected_effects)
        self.assertEqual(actual[requests[-1]], [])

        opened = []
        def empty_connect():
            connection = sqlite3.connect(":memory:", factory=TrackingConnection)
            connection.row_factory = sqlite3.Row
            opened.append(connection)
            return connection
        with patch.object(repository, "connect", side_effect=empty_connect):
            with self.assertRaises(sqlite3.OperationalError):
                repository.fetch_ranked_effects_many([real_key])
            with self.assertRaises(sqlite3.OperationalError):
                repository.list_mods_for_slot("Rifle")
        self.assertTrue(all(connection.closed for connection in opened))

    def test_real_exclusive_lock_is_not_reported_as_empty_catalog_or_effects(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "locked.db"
            with closing(sqlite3.connect(DATABASE)) as source, \
                 closing(sqlite3.connect(path)) as target:
                source.backup(target)
            mods = ModRepository(path)
            weapons = WeaponRepository(path)
            mod = mods.list_mods_for_slot("PRIMARY")[0]

            def short_connect():
                connection = sqlite3.connect(path, timeout=0.01, factory=TrackingConnection)
                connection.row_factory = sqlite3.Row
                return connection

            with closing(sqlite3.connect(path)) as locker, \
                 patch.object(mods, "connect", side_effect=short_connect), \
                 patch.object(weapons, "connect", side_effect=short_connect):
                locker.execute("BEGIN EXCLUSIVE")
                for query in (
                    lambda: mods.list_mods_for_slot("PRIMARY"),
                    lambda: mods.get_mod(mod.unique_name),
                    lambda: mods.fetch_ranked_effects(mod.unique_name, 0),
                    lambda: mods.fetch_ranked_effects_many([(mod.unique_name, 0)]),
                    lambda: weapons.search_weapons("Hek"),
                    lambda: weapons.get_weapon("Hek"),
                ):
                    with self.assertRaisesRegex(sqlite3.OperationalError, "locked"):
                        query()
                locker.rollback()
                self.assertIsNotNone(weapons.get_weapon("Hek"))
                self.assertTrue(mods.list_mods_for_slot("PRIMARY"))

    def test_database_rebuild_closes_and_rolls_back_on_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "weapons.db"
            payload = ImportPayload([], [], [], [])
            connections = []
            real_connect = sqlite3.connect
            def connect(*args, **kwargs):
                connection = real_connect(*args, factory=TrackingConnection, **kwargs)
                connections.append(connection)
                return connection
            def fail(connection, payload):
                connection.execute("CREATE TABLE should_rollback (n INTEGER)")
                raise RuntimeError("simulated import failure")
            with patch("modules.importer.load_import_payload", return_value=payload), \
                 patch("modules.importer.sqlite3.connect", side_effect=connect), \
                 patch("modules.importer.apply_import_payload", side_effect=fail):
                with self.assertRaisesRegex(RuntimeError, "simulated import failure"):
                    build_database(path, "weapons.zip", "mods.zip")
            self.assertTrue(connections[0].closed)
            with closing(real_connect(path)) as connection:
                self.assertIsNone(connection.execute(
                    "SELECT name FROM sqlite_master WHERE name='should_rollback'").fetchone())

    def test_spreadsheet_is_closed_if_worksheet_is_missing(self):
        class Workbook:
            closed = False
            def __getitem__(self, key):
                raise KeyError(key)
            def close(self):
                self.closed = True
        workbook = Workbook()
        with patch("modules.importer.openpyxl.load_workbook", return_value=workbook):
            with self.assertRaises(KeyError):
                rows_from_spreadsheet("missing-sheet.xlsx")
        self.assertTrue(workbook.closed)

    def test_spreadsheet_rows_are_cached_but_never_stale(self):
        import tempfile
        from pathlib import Path

        def write(path, name):
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Weapon-Data"
            sheet.append(["Weapon", "Total Damage"] + [""] * 19)
            sheet.append(["header row"])
            sheet.append([name, 100])
            workbook.save(path)

        with tempfile.TemporaryDirectory() as tmp:
            book = Path(tmp) / "weapons.xlsx"
            write(book, "First Weapon")
            first = rows_from_spreadsheet(book)
            self.assertEqual([row["weapon"] for row in first if row["weapon"]],
                             ["First Weapon"])
            # Cached reads are equal but independent objects.
            second = rows_from_spreadsheet(book)
            self.assertEqual(first, second)
            self.assertIsNot(first[0], second[0])
            second[0]["weapon"] = "MUTATED"
            self.assertEqual(rows_from_spreadsheet(book)[0]["weapon"],
                             first[0]["weapon"])
            # A rewritten workbook invalidates the cached parse.
            write(book, "Renamed Weapon Two")
            self.assertEqual([row["weapon"] for row in rows_from_spreadsheet(book)
                              if row["weapon"]], ["Renamed Weapon Two"])

if __name__ == "__main__":
    unittest.main()
